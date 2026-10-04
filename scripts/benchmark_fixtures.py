#!/usr/bin/env python3
"""Generate small, deterministic documents for the cross-host benchmark.

Run inside the Dethrottled API image, which already contains the document
libraries. The output directory can then be served by a plain HTTP server.
"""
from __future__ import annotations

import argparse
import io
from pathlib import Path

import pymupdf as fitz
from docx import Document
from ebooklib import epub
from odf import text as odf_text
from odf.opendocument import OpenDocumentText
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

MARKER = "Dethrottled benchmark: cedar observatory measured 5120 megawatts in 2026."


def make(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    prose = " ".join([MARKER, "The source table links generation, storage, and demand."] * 70)
    paras = "\n".join(f"<p>{MARKER} Section {i}. {prose[:300]}</p>" for i in range(35))
    (out / "article.html").write_text(
        f"<!doctype html><title>Cedar energy report</title><main><h1>Cedar energy report</h1>"
        f"<article>{paras}</article></main>", encoding="utf-8")
    (out / "table.html").write_text(
        "<!doctype html><title>Capacity table</title><main><h1>Capacity table</h1>"
        + "<table>" + "".join(
            f"<tr><td>Station {i}</td><td>{5120 + i}</td><td>{MARKER}</td></tr>"
            for i in range(120)) + "</table></main>", encoding="utf-8")
    (out / "values.csv").write_text(
        "station,year,megawatts,note\n" + "".join(
            f"Cedar {i},2026,{5120+i},{MARKER}\n" for i in range(500)), encoding="utf-8")
    (out / "wrong.pdf").write_text(
        "<!doctype html><title>404 masquerading as PDF</title><p>Document not found.</p>",
        encoding="utf-8")

    pdf = fitz.open()
    for i in range(3):
        page = pdf.new_page()
        page.insert_text((50, 70), f"Cedar energy report page {i + 1}", fontsize=17)
        for line in range(24):
            page.insert_text((50, 100 + line * 20), f"{line + 1}. {MARKER}", fontsize=10)
    pdf.save(out / "report.pdf")
    pdf.close()

    image = Image.new("RGB", (1800, 700), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 38)
    for i in range(5):
        draw.text((40, 40 + i * 110), MARKER, fill="black", font=font)
    scan = fitz.open()
    page = scan.new_page(width=900, height=350)
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    page.insert_image(page.rect, stream=stream.getvalue())
    scan.save(out / "scan.pdf")
    scan.close()

    doc = Document()
    doc.add_heading("Cedar energy report", 0)
    for i in range(20):
        doc.add_paragraph(f"Section {i}. {MARKER} {prose[:250]}")
    doc.save(out / "report.docx")

    book = Workbook()
    sheet = book.active
    sheet.append(("station", "year", "megawatts", "note"))
    for i in range(500):
        sheet.append((f"Cedar {i}", 2026, 5120 + i, MARKER))
    book.save(out / "report.xlsx")

    deck = Presentation()
    for i in range(8):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
        box.text = f"Cedar capacity slide {i + 1}. {MARKER} " * 3
    deck.save(out / "report.pptx")

    odt = OpenDocumentText()
    for i in range(20):
        odt.text.addElement(odf_text.P(text=f"Section {i}. {MARKER} {prose[:250]}"))
    odt.save(str(out / "report.odt"))

    book = epub.EpubBook()
    book.set_identifier("dethrottled-benchmark-cedar")
    book.set_title("Cedar energy report")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Report", file_name="report.xhtml", lang="en")
    chapter.content = f"<h1>Cedar energy report</h1>{paras}"
    book.add_item(chapter)
    book.toc = (epub.Link("report.xhtml", "Report", "report"),)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(str(out / "report.epub"), book)

    (out / "report.rtf").write_text(
        "{\\rtf1\\ansi\\deff0 " + "\\par ".join([MARKER] * 30) + "}",
        encoding="ascii")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("out", type=Path)
    make(parser.parse_args().out)
