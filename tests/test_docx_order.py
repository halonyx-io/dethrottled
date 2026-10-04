"""Narrative around a table should remain attached to the table it explains."""

import io

import docx

from dethrottled import documents


def test_docx_paragraphs_and_tables_follow_source_order():
    document = docx.Document()
    document.add_paragraph("Fiscal year 2025")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Country"
    table.cell(0, 1).text = "Megawatts"
    table.cell(1, 0).text = "Denmark"
    table.cell(1, 1).text = "5120"
    document.add_paragraph("Excludes offshore capacity")
    buffer = io.BytesIO()
    document.save(buffer)

    text, reason = documents.to_text(buffer.getvalue(), "docx", 20000)
    assert reason == ""
    assert text.splitlines() == [
        "Fiscal year 2025", "Country | Megawatts", "Denmark | 5120",
        "Excludes offshore capacity"]
